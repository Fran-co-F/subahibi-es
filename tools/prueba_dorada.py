# SPDX-License-Identifier: GPL-3.0-or-later
"""Prueba dorada: corre el flujo completo sobre el juego real y guarda una huella de todo lo que genera.

    python tools/prueba_dorada.py JUEGO SALIDA [--parche-viejo ZIP] [--referencia OTRA_SALIDA]

Arma en SALIDA/juego una copia de la instalación hecha con enlaces (el juego original no se toca), y ahí corre
init -> exportar -> (copia 2 originales a imagenes/traducidas) -> exportar -> build -> build --jugable -> estado -vv
con el código de este repositorio (src/).  Escribe en SALIDA:
    hashes.txt   sha256 de cada archivo generado (salida/, originales/, imagenes/, jugable/ sin los enlaces)
    texto.txt    el contenido de todas las hojas, fila por fila
    registro.txt lo que imprimió cada comando
Con --referencia compara hashes.txt y texto.txt con los de otra corrida (por ejemplo, la de main) y sale con 1 si difieren.
Para obtener la referencia de main sin cambiar de rama:  git worktree add /tmp/main main  y correr su tools/prueba_dorada.py.
"""
import argparse
import glob
import hashlib
import os
import shutil
import subprocess
import sys

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _cli(cwd, log, *args):
    env = dict(os.environ, PYTHONPATH=os.path.join(REPO, "src"))
    for e in ("TL_ORGANIZER_PROJECT", "SUBAHIBI_ES_PROJECT"):
        env.pop(e, None)
    r = subprocess.run([sys.executable, "-m", "tl_organizer", *args], cwd=cwd, env=env, capture_output=True, text=True)
    log.write(f"$ tl-organizer {' '.join(args)}\n{r.stdout}{r.stderr}\n")
    if r.returncode:
        sys.exit(f"falló «tl-organizer {' '.join(args)}»:\n{r.stdout}{r.stderr}")


def _enlazar_juego(src, dst):
    os.makedirs(dst)
    for f in sorted(os.listdir(src)):
        p = os.path.join(src, f)
        if os.path.isfile(p) and f not in ("tl-organizer.json", "subahibi-es.json"):
            try:
                os.symlink(p, os.path.join(dst, f))
            except OSError:
                shutil.copy2(p, os.path.join(dst, f))


def _hashes(work):
    out = []
    for top in ("salida", "originales", "imagenes", "jugable"):
        for root, _d, files in os.walk(os.path.join(work, top)):
            for f in files:
                p = os.path.join(root, f)
                if os.path.islink(p):
                    continue                      # los enlaces de --jugable apuntan al juego
                with open(p, "rb") as fh:
                    out.append(f"{hashlib.sha256(fh.read()).hexdigest()}  {os.path.relpath(p, work).replace(os.sep, '/')}")
    return sorted(out, key=lambda l: l[66:])


def _texto(work):
    from openpyxl import load_workbook
    out = []
    hojas = sorted(glob.glob(os.path.join(work, "hojas", "*.xlsx"))) or sorted(glob.glob(os.path.join(work, "texto", "*.xlsx")))
    for p in hojas:                               # (texto/: versiones anteriores a traduccion/ + hojas/)
        wb = load_workbook(p, read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append(f"{os.path.basename(p)} {ws.title} {row[:8]!r}")     # sin la columna oculta «base»
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("juego"); ap.add_argument("salida")
    ap.add_argument("--parche-viejo"); ap.add_argument("--referencia")
    a = ap.parse_args()
    if os.path.exists(a.salida):
        sys.exit(f"{a.salida} ya existe: elegí otra carpeta (no se borra nada)")
    juego = os.path.join(a.salida, "juego")
    _enlazar_juego(os.path.abspath(a.juego), juego)
    with open(os.path.join(a.salida, "registro.txt"), "w", encoding="utf-8") as log:
        _cli(juego, log, "init", *(["--parche-viejo", os.path.abspath(a.parche_viejo)] if a.parche_viejo else []))
        _cli(juego, log, "exportar")
        work = next(w for w in (os.path.join(juego, "tl-organizer-trabajo"), os.path.join(juego, "subahibi-es-trabajo"))
                    if os.path.isdir(w))      # (la segunda: versiones anteriores al cambio de nombre)
        for png in sorted(glob.glob(os.path.join(work, "originales", "*.png")))[:2]:
            shutil.copy2(png, os.path.join(work, "imagenes", "traducidas"))
        _cli(juego, log, "exportar")
        _cli(juego, log, "build")
        _cli(juego, log, "build", "--jugable")
        _cli(juego, log, "estado", "-vv")
    res = {"hashes.txt": _hashes(work), "texto.txt": _texto(work)}
    for name, lines in res.items():
        with open(os.path.join(a.salida, name), "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")
    print(f"{len(res['hashes.txt'])} archivos, {len(res['texto.txt'])} filas de hojas -> {a.salida}")
    if a.referencia:
        bad = False
        for name, lines in res.items():
            with open(os.path.join(a.referencia, name), encoding="utf-8") as f:
                ref = f.read().splitlines()
            if ref == lines:
                print(f"{name}: idéntico")
            else:
                bad = True
                faltan, sobran = set(ref) - set(lines), set(lines) - set(ref)
                print(f"{name}: DIFIERE ({len(faltan)} líneas solo en la referencia, {len(sobran)} solo en esta corrida)")
                for l in sorted(faltan)[:5]:
                    print("  -", l[:160])
                for l in sorted(sobran)[:5]:
                    print("  +", l[:160])
        return 1 if bad else 0
    return 0


if __name__ == "__main__":
    sys.exit(main())
