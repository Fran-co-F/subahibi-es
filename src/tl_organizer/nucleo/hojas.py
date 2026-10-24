# SPDX-License-Identifier: GPL-3.0-or-later
"""Hojas de traducción: una fila por texto.  Formatos: TSV (por script, estable para git) y XLSX (por .arc,
una hoja por script).

Columnas: clave, tipo, <idioma original>, <idioma de la traducción>, estado, nota, pos, hash
  En las filas (dicts) las columnas del texto se llaman siempre «original» y «traduccion»; en la hoja, el encabezado lleva el
  nombre de cada idioma del proyecto (con japonés -> español: «japones» y «espanol»).  Al leer se acepta cualquier nombre en
  esas dos columnas.
  clave   = <script>:<número de fila>        pos  = desplazamiento del token dentro del escenario
  hash    = sha1 corto del texto japonés (avisa si el juego cambió)
  estado  = sin traducir | traducido | revisar | del parche
En XLSX los saltos de línea son saltos reales de celda; en TSV se escriben como \\n (y \\t, \\\\).
"""
import hashlib

COLUMNS = ["clave", "tipo", "original", "traduccion", "estado", "nota", "pos", "hash"]     # claves de cada fila
ESTADOS = ["sin traducir", "traducido", "revisar", "del parche"]
IDIOMAS = ("ja", "es")                    # (original, traducción) si el proyecto no dice otra cosa
NOMBRES_IDIOMA = {"ja": "japones", "es": "espanol", "en": "ingles", "zh": "chino", "ko": "coreano", "fr": "frances",
                  "de": "aleman", "it": "italiano", "pt": "portugues", "ru": "ruso", "ca": "catalan", "gl": "gallego"}


def nombre_idioma(codigo):
    """Nombre de la columna para un idioma (sin tildes, como los demás encabezados); un código desconocido se usa tal cual."""
    return NOMBRES_IDIOMA.get(codigo, codigo)


def encabezado(idiomas=IDIOMAS):
    return ["clave", "tipo", nombre_idioma(idiomas[0]), nombre_idioma(idiomas[1]), "estado", "nota", "pos", "hash"]


def _check_head(head, where):
    if len(head) < len(COLUMNS) or head[:2] != COLUMNS[:2] or head[4:8] != COLUMNS[4:8] or not head[2] or not head[3]:
        raise ValueError(f"{where}: cabecera inesperada {head}")


def jp_hash(raw):
    return hashlib.sha1(raw).hexdigest()[:8]


def _esc(s):
    return s.replace("\\", "\\\\").replace("\t", "\\t").replace("\r", "").replace("\n", "\\n")


def _unesc(s):
    out, i = [], 0
    while i < len(s):
        c = s[i]
        if c == "\\" and i + 1 < len(s):
            n = s[i + 1]
            out.append({"n": "\n", "t": "\t", "\\": "\\"}.get(n, "\\" + n)); i += 2
        else:
            out.append(c); i += 1
    return "".join(out)


def write_tsv(path, rows, idiomas=IDIOMAS):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\t".join(encabezado(idiomas)) + "\n")
        for r in rows:
            f.write("\t".join(_esc(str(r.get(c, ""))) for c in COLUMNS) + "\n")


def read_tsv(path):
    with open(path, encoding="utf-8") as f:
        lines = f.read().split("\n")
    _check_head(lines[0].split("\t"), path)
    rows = []
    for ln in lines[1:]:
        if not ln:
            continue
        p = ln.split("\t")
        p += [""] * (len(COLUMNS) - len(p))
        rows.append({c: _unesc(v) for c, v in zip(COLUMNS, p)})
    return rows


def write_xlsx(path, sheets, idiomas=IDIOMAS):
    """sheets: {nombre_de_hoja: [filas]}"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    wb.remove(wb.active)
    widths = {"A": 22, "B": 11, "C": 60, "D": 60, "E": 13, "F": 24, "G": 9, "H": 10}
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        ws.append(encabezado(idiomas))
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2E4057")
        for r in rows:
            ws.append([r.get(c, "") if c != "pos" else r.get(c, "") for c in COLUMNS])
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        wrap = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2, max_col=len(COLUMNS)):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        if rows:
            dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"E2:E{len(rows) + 1}")
    wb.save(path)


def read_xlsx(path):
    """-> {hoja: [filas]}"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        head = [str(x) if x is not None else "" for x in next(it, [])]
        _check_head(head, f"{path}[{ws.title}]")
        rows = []
        for vals in it:
            if vals is None or all(v is None for v in vals):
                continue
            vals = list(vals) + [None] * (len(COLUMNS) - len(vals))
            rows.append({c: ("" if v is None else str(v)) for c, v in zip(COLUMNS, vals)})
        out[ws.title] = rows
    return out


def safe_name(s):
    return "".join("_" if ch in '[]:*?/\\' else ch for ch in s)
